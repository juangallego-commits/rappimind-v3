# -*- coding: utf-8 -*-
"""
Arma el libro de Excel de la revision legal global de RappiMind.

No es parte del build ni de las pruebas: es una utilidad de una sola vez para
preparar el Google Sheet donde los equipos legales de los nueve paises dejan su
feedback. Se sube a Drive y Google lo convierte en Sheet, con sus pestanas, sus
listas desplegables y sus formulas.

    pip install openpyxl
    python3 scripts/hoja_feedback.py <ruta_de_salida.xlsx>

Los datos se leen del repositorio, no se escriben a mano: data/paises.json para
la pestana de datos legales y modelos/leyendas/ para la de leyendas. Si manana
cambia un valor, se vuelve a correr y la hoja sale al dia.
"""

import json
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FECHA_LIMITE = 'viernes 25 de septiembre de 2026, 6:00 p.m. hora de Colombia'

# ---------------------------------------------------------------- estilos ----

NARANJA = 'FF441F'
NEGRO = '16181D'
GRIS = 'F4F5F7'
GRIS_BORDE = 'E3E5EA'
AMARILLO = 'FFF6E5'
VERDE = 'EAF6EE'

BORDE = Border(*[Side(style='thin', color=GRIS_BORDE)] * 4)

TITULO = Font(name='Arial', size=14, bold=True, color='FFFFFF')
CABECERA = Font(name='Arial', size=10, bold=True, color='FFFFFF')
CABECERA_JUAN = Font(name='Arial', size=10, bold=True, color=NEGRO)
NORMAL = Font(name='Arial', size=10)
NOTA = Font(name='Arial', size=10, italic=True, color='5D626D')
FUERTE = Font(name='Arial', size=10, bold=True)

FONDO_TITULO = PatternFill('solid', fgColor=NEGRO)
FONDO_CABECERA = PatternFill('solid', fgColor=NARANJA)
FONDO_CABECERA_JUAN = PatternFill('solid', fgColor=GRIS_BORDE)
FONDO_NOTA = PatternFill('solid', fgColor=GRIS)
FONDO_EJEMPLO = PatternFill('solid', fgColor=AMARILLO)
FONDO_PREFIJADO = PatternFill('solid', fgColor=VERDE)

ARRIBA = Alignment(vertical='top', wrap_text=True)
ARRIBA_IZQ = Alignment(vertical='top', horizontal='left', wrap_text=True)


def leer_paises():
    ruta = os.path.join(RAIZ, 'data', 'paises.json')
    with open(ruta, encoding='utf-8') as archivo:
        return json.load(archivo)['paises']


def leer_leyendas():
    carpeta = os.path.join(RAIZ, 'modelos', 'leyendas')
    leyendas = []
    for nombre in sorted(os.listdir(carpeta)):
        if not nombre.endswith('.md'):
            continue
        with open(os.path.join(carpeta, nombre), encoding='utf-8') as archivo:
            crudo = archivo.read()
        meta, cuerpo = {}, []
        for linea in crudo.splitlines():
            marca = re.match(r'^%%\s*(\w+):\s*(.*)$', linea)
            if marca:
                meta[marca.group(1)] = marca.group(2).strip()
            elif linea.strip():
                cuerpo.append(linea.strip())
        meta['texto'] = ' '.join(cuerpo)
        meta['archivo'] = nombre
        leyendas.append(meta)
    return leyendas


NOMBRES_TIPOS = {
    'cashback': 'Cashback',
    'cashback_day': 'Cashback Multitienda',
    'carrera_compras': 'Carrera de Compras',
    'cashback_travel': 'Cashback Travel',
    'descuento_travel': 'Descuento Travel',
    'cashback_eta': '15 minutos o Cashback',
    'promocao_br': 'Promoção Brasil',
}

NOMBRES_PAISES = {
    'CO': 'Colombia', 'MX': 'México', 'PE': 'Perú', 'AR': 'Argentina',
    'CL': 'Chile', 'EC': 'Ecuador', 'UY': 'Uruguay', 'CR': 'Costa Rica',
    'BR': 'Brasil',
}


def paises_en_letra(codigos):
    """Convierte "CL, CR, UY" en "Chile, Costa Rica, Uruguay"."""
    partes = [p.strip() for p in (codigos or '').split(',') if p.strip()]
    if len(partes) > 4:
        return 'Aplica a todos'
    return ', '.join(NOMBRES_PAISES.get(p, p) for p in partes)


DONDE_SE_ELIGE = {
    'pregunta': 'La herramienta la ofrece y el comercial la elige',
    'condicion_turbo': 'Sale sola cuando la campaña es de Turbo',
    'no_se_ofrece': 'Hoy no se ofrece: la categoría está cerrada en ese país',
    'ninguna_parte_todavia': 'Está escrita pero todavía no se ofrece en ninguna pregunta',
}


def hoja_nueva(libro, nombre, titulo, instruccion, cabeceras, anchos, corte='A4'):
    """Crea una pestaña con su banda de título, su instrucción y sus cabeceras."""
    hoja = libro.create_sheet(nombre)
    ultima = get_column_letter(len(cabeceras))

    hoja.merge_cells('A1:%s1' % ultima)
    celda = hoja['A1']
    celda.value = titulo
    celda.font = TITULO
    celda.fill = FONDO_TITULO
    celda.alignment = Alignment(vertical='center', horizontal='left')
    hoja.row_dimensions[1].height = 30

    hoja.merge_cells('A2:%s2' % ultima)
    celda = hoja['A2']
    celda.value = instruccion
    celda.font = NOTA
    celda.fill = FONDO_NOTA
    celda.alignment = ARRIBA_IZQ
    hoja.row_dimensions[2].height = 34

    for indice, texto in enumerate(cabeceras, start=1):
        celda = hoja.cell(row=3, column=indice, value=texto)
        de_juan = texto.startswith('[')
        celda.font = CABECERA_JUAN if de_juan else CABECERA
        celda.fill = FONDO_CABECERA_JUAN if de_juan else FONDO_CABECERA
        celda.alignment = Alignment(vertical='center', wrap_text=True)
        celda.border = BORDE
    hoja.row_dimensions[3].height = 42

    for indice, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = ancho

    hoja.freeze_panes = corte
    hoja.auto_filter.ref = 'A3:%s3' % ultima
    return hoja


def escribir_fila(hoja, fila, valores, fondo=None, fuente=None):
    for indice, valor in enumerate(valores, start=1):
        celda = hoja.cell(row=fila, column=indice, value=valor)
        celda.font = fuente or NORMAL
        celda.alignment = ARRIBA
        celda.border = BORDE
        if fondo is not None:
            celda.fill = fondo
    return fila + 1


def validacion(hoja, rango_listas, rango_celdas):
    regla = DataValidation(type='list', formula1=rango_listas, allow_blank=True)
    regla.error = 'Elige un valor de la lista.'
    regla.errorTitle = 'Valor fuera de la lista'
    hoja.add_data_validation(regla)
    regla.add(rango_celdas)


# ------------------------------------------------------------------ LEEME ----

LEEME = [
    ('h1', 'RappiMind · Revisión legal global'),
    ('p', 'Léelo una vez. Toma tres minutos y te ahorra media hora.'),

    ('h2', 'Qué es RappiMind'),
    ('li', 'Es una aplicación interna donde el comercial, el CAM y el KAM, genera solo los Términos y Condiciones de su campaña de marketing, sin escribirle a Legal.'),
    ('li', 'El comercial solo responde hechos: la tienda, las fechas, el porcentaje, el tope, el presupuesto. El texto legal sale completo de los modelos que ya aprobamos nosotros, palabra por palabra.'),
    ('li', 'Si falta un dato, no genera. Nunca sale un documento con espacios en blanco, con marcadores sin llenar ni con frases a medias.'),
    ('li', 'Hoy cubre nueve países, siete tipos de campaña y cincuenta y cinco leyendas de piezas publicitarias.'),

    ('h2', 'Qué estamos probando y hasta cuándo'),
    ('li', 'Son pruebas internas, solo entre abogados. Ningún comercial tiene la herramienta todavía y nada de lo que generes aquí se publica.'),
    ('li', 'La fecha para dejar comentarios es el %s.' % FECHA_LIMITE),
    ('li', 'Lo que llegue después del viernes entra en la siguiente ronda, no en esta.'),

    ('h2', 'Antes de escribir en esta hoja'),
    ('p', 'Entra a la herramienta y genera dos o tres campañas de tu país, con datos inventados pero realistas. El enlace está en el mensaje del canal. Después lee el documento completo, de principio a fin. Casi todo lo que vale la pena encontrar aparece leyendo el documento entero, no la pantalla.'),

    ('h2', 'Lo mínimo que necesito de ti, entre 30 y 45 minutos'),
    ('li', '1. Pestaña "2 Hallazgos": todo lo que esté mal. Una fila por hallazgo.'),
    ('li', '2. Pestaña "3 Datos de tu país": confirma o corrige los datos legales de tu país. Esos datos entran en cada documento que se genere allá.'),
    ('li', '3. Pestaña "6 Accesos Squarespace": dime qué persona de tu país necesita usuario para publicar los T&C.'),

    ('h2', 'Si te queda tiempo'),
    ('li', '4. Pestaña "4 Preguntas de Legal": diez preguntas de fondo sobre la ley de tu país.'),
    ('li', '5. Pestaña "5 Leyendas": las advertencias obligatorias que van en las piezas publicitarias.'),
    ('li', '6. Pestaña "7 Accesos RappiMind": quién más de tu equipo quiere probar la herramienta.'),
    ('li', '7. Pestaña "8 Cobertura": qué campaña te piden mucho y la herramienta todavía no hace.'),

    ('h2', 'Cinco reglas para que esto sirva'),
    ('li', '1. Una fila por hallazgo. Si son tres cosas distintas dentro de la misma cláusula, son tres filas.'),
    ('li', '2. Pega el texto tal como sale hoy y escribe al lado el texto exacto que debería salir. "Hay que mejorar la redacción" no lo puedo aplicar a nada.'),
    ('li', '3. Filtra por tu país. Las pestañas largas ya traen las filas de los nueve países.'),
    ('li', '4. No borres filas ni columnas, ni cambies las cabeceras. Esta hoja la leo con un script.'),
    ('li', '5. Las columnas grises que empiezan con [Juan] las lleno yo. No escribas ahí.'),

    ('h2', 'Cómo clasificamos la severidad'),
    ('li', 'Bloqueante: publicado así nos expone a una sanción, o el documento dice algo que no es cierto. No se puede sacar.'),
    ('li', 'Serio: está mal y hay que arreglarlo antes de abrirle la herramienta a comercial, pero no nos expone hoy.'),
    ('li', 'Menor: erratas, tildes, una palabra que en tu país no se dice así.'),
    ('li', 'Sugerencia: una mejora. No es un error.'),

    ('h2', 'Qué NO hace la herramienta hoy, para que no lo busques'),
    ('li', 'No publica sola en Squarespace. Genera el documento y una persona lo pega en la página.'),
    ('li', 'No cubre sorteos, rifas, raspa y gana ni nada con azar. Eso sigue llegando a Legal y la herramienta lo desvía sola.'),
    ('li', 'No deja editar el documento. El comercial no escribe ni una coma de texto legal, y eso es a propósito.'),
    ('li', 'No manda nada a ninguna autoridad ni hace registros previos de promociones.'),

    ('h2', 'La regla que no se negocia'),
    ('p', 'Rappi nunca es el Organizador de la campaña. El Organizador es siempre la tienda aliada o la marca que fondea el beneficio. Rappi es plataforma de intermediación y responsable del tratamiento de los datos personales. Que Rappi opere la mecánica, cargue los Créditos o gestione la entrega de un premio no la convierte en organizadora. La única excepción autorizada es el modelo de Brasil, que está escrito así desde sus T&C aprobados, y no sirve de precedente para ningún otro país. Si tu hallazgo choca con esta regla, escríbelo igual y lo discutimos.'),

    ('h2', 'Qué pasa después del viernes'),
    ('li', 'El lunes leo todo, marco cada fila con mi decisión en las columnas grises y respondo por el canal.'),
    ('li', 'Lo aceptado entra a los modelos y vuelvo a generar los documentos de tu país.'),
    ('li', 'Te aviso qué quedó y qué no, con la razón. Nada se queda sin respuesta.'),

    ('h2', 'Dónde preguntar'),
    ('p', 'En el hilo del canal #legalglobal. Dueño del producto y aprobador legal: Juan Camilo Gallego, Legal Colombia, juan.gallego@rappi.com.'),
]


def pestana_leeme(libro):
    hoja = libro.create_sheet('LÉEME')
    hoja.column_dimensions['A'].width = 4
    hoja.column_dimensions['B'].width = 118
    hoja.sheet_view.showGridLines = False

    fila = 1
    for clase, texto in LEEME:
        celda = hoja.cell(row=fila, column=2, value=texto)
        celda.alignment = ARRIBA_IZQ
        if clase == 'h1':
            celda.font = Font(name='Arial', size=16, bold=True, color='FFFFFF')
            celda.fill = FONDO_TITULO
            hoja.cell(row=fila, column=1).fill = FONDO_TITULO
            hoja.row_dimensions[fila].height = 34
        elif clase == 'h2':
            celda.font = Font(name='Arial', size=12, bold=True, color=NARANJA)
            hoja.row_dimensions[fila].height = 30
        elif clase == 'li':
            celda.font = NORMAL
            celda.value = '•  ' + texto
            hoja.row_dimensions[fila].height = max(16, 15 * (len(texto) // 105 + 1))
        else:
            celda.font = NORMAL
            hoja.row_dimensions[fila].height = max(16, 15 * (len(texto) // 105 + 1))
        fila += 1
    return hoja


# ----------------------------------------------------------------- listas ----

LISTAS = {
    'paises': ['Colombia', 'México', 'Perú', 'Argentina', 'Chile', 'Ecuador',
               'Uruguay', 'Costa Rica', 'Brasil', 'Aplica a todos'],
    'tipos': ['Cashback', 'Cashback Multitienda', 'Carrera de Compras',
              'Cashback Travel', 'Descuento Travel', '15 minutos o Cashback',
              'Promoção Brasil', 'T&C generales de descuento', 'Aplica a todos',
              'No aplica'],
    'donde': ['Documento generado', 'Pregunta del formulario',
              'Pantalla de resumen', 'Leyenda de la pieza', 'Correo que llega',
              'Historial', 'Ayuda', 'Datos de mi país', 'Traducción al portugués',
              'Otra cosa'],
    'severidad': ['Bloqueante', 'Serio', 'Menor', 'Sugerencia'],
    'si_no': ['Sí', 'No', 'No sé'],
    'estado_juan': ['Nuevo', 'Leído', 'Aceptado', 'Aplicado', 'No se aplica',
                    'Necesito hablarlo'],
    'esta_bien': ['Está bien', 'Hay que cambiarlo', 'No sé', 'No aplica en mi país'],
    'respuesta': ['Sí', 'No', 'Depende', 'No sé'],
    'prioridad': ['Alta', 'Media', 'Baja'],
    'estado_acceso': ['Pendiente', 'Pedido a Maca', 'Creado', 'Ya lo tenía', 'Rechazado'],
    'rol_rappimind': ['legal', 'comercial', 'admin'],
    'para_que': ['Publicar T&C de cashback', 'Publicar T&C de concursos',
                 'Publicar los dos', 'Solo consultar'],
    'avance': ['Sin empezar', 'En eso', 'Terminado', 'No voy a alcanzar'],
    'estado_alta': ['Pendiente', 'Dado de alta', 'No aplica'],
}


def pestana_listas(libro):
    hoja = libro.create_sheet('Listas')
    for columna, (nombre, valores) in enumerate(LISTAS.items(), start=1):
        celda = hoja.cell(row=1, column=columna, value=nombre)
        celda.font = FUERTE
        for fila, valor in enumerate(valores, start=2):
            hoja.cell(row=fila, column=columna, value=valor).font = NORMAL
        hoja.column_dimensions[get_column_letter(columna)].width = 26
    hoja.sheet_state = 'hidden'
    return hoja


def rango(nombre):
    """El rango de una lista, listo para una validación."""
    columna = get_column_letter(list(LISTAS).index(nombre) + 1)
    return '=Listas!$%s$2:$%s$%d' % (columna, columna, len(LISTAS[nombre]) + 1)


# -------------------------------------------------------------- checklist ----

TAREAS = [
    '1. Generé 2 o 3 campañas de mi país',
    '2. Leí el documento completo, de principio a fin',
    '3. Dejé mis hallazgos en la pestaña 2',
    '4. Revisé los datos legales de mi país en la pestaña 3',
    '5. Dije quién necesita acceso a Squarespace en la pestaña 6',
]


def pestana_checklist(libro, paises):
    cabeceras = (['País', 'Abogado que revisa', 'Correo'] + TAREAS +
                 ['Cómo voy', 'Fecha en que terminé', 'Lo que quiero comentar de viva voz',
                  '[Juan] Nota'])
    anchos = [14, 24, 26, 22, 22, 22, 22, 22, 16, 16, 40, 30]
    hoja = hoja_nueva(
        libro, '1 Checklist',
        '1. Checklist por país',
        'Busca la fila de tu país, escribe tu nombre y ve marcando. Es para que los dos sepamos '
        'quién va cerrando y a quién le falta. Fecha límite: %s.' % FECHA_LIMITE,
        cabeceras, anchos)

    fila = 4
    for pais in paises:
        fila = escribir_fila(hoja, fila, [pais['nombre'], '', '', '', '', '', '', '', 'Sin empezar'])
    for columna in range(4, 9):
        letra = get_column_letter(columna)
        validacion(hoja, rango('si_no'), '%s4:%s%d' % (letra, letra, fila - 1))
    validacion(hoja, rango('avance'), 'I4:I%d' % (fila - 1))
    return hoja


# -------------------------------------------------------------- hallazgos ----

EJEMPLOS_HALLAZGOS = [
    ['EJEMPLO, no lo borres', 'Ana Ejemplo', 'ana@rappi.com', 'Chile', 'Cashback',
     'Documento generado', 'Cláusula 7, Ley aplicable',
     'Para todo lo relacionado con la interpretación de estos Términos se aplicarán las leyes '
     'de la República de Chile.',
     'Estos Términos se rigen por las leyes de la República de Chile y en especial por la Ley '
     '19.496 sobre protección de los derechos de los consumidores.',
     'La Ley 19.496 es la norma de consumo y el SERNAC exige que se cite cuando la promoción se '
     'dirige a consumidores finales.',
     'Serio', 'No', ''],
    ['EJEMPLO, no lo borres', 'Ana Ejemplo', 'ana@rappi.com', 'Chile', 'Cashback',
     'Pregunta del formulario', 'Pregunta del tope por pedido',
     'La pregunta deja poner un tope mayor que el presupuesto total y no avisa nada.',
     'Que avise en pantalla cuando el tope sea mayor que el presupuesto, porque entonces el '
     'primer pedido se come la campaña entera.',
     'No es un tema de ley, es que el documento queda diciendo algo que no se puede cumplir.',
     'Menor', 'No', ''],
]


def pestana_hallazgos(libro, filas_libres=400):
    cabeceras = [
        'Fecha', 'Quién reporta', 'Correo', 'País', 'Tipo de campaña',
        'Dónde lo viste', 'Cláusula, sección o pantalla',
        'Qué dice hoy (pega el texto tal cual)',
        'Qué debería decir (el texto exacto)',
        'Por qué: norma, riesgo o razón',
        'Severidad', '¿Bloquea publicar?',
        'Enlace al documento que generaste',
        '[Juan] Estado', '[Juan] Decisión', '[Juan] Dónde se aplica', '[Juan] Fecha',
    ]
    anchos = [14, 20, 26, 14, 20, 22, 26, 52, 52, 40, 14, 14, 30, 14, 42, 26, 12]
    hoja = hoja_nueva(
        libro, '2 Hallazgos',
        '2. Hallazgos: todo lo que esté mal',
        'Una fila por hallazgo. Pega el texto que sale hoy y escribe al lado el texto exacto que '
        'debería salir. Las dos primeras filas son ejemplos, déjalas donde están. Las columnas '
        'grises las lleno yo.',
        cabeceras, anchos)

    fila = 4
    for ejemplo in EJEMPLOS_HALLAZGOS:
        fila = escribir_fila(hoja, fila, ejemplo + ['Nuevo'], fondo=FONDO_EJEMPLO)

    # Las filas vacias no se pintan una por una: el archivo se hincha y Google
    # les pone su propio formato al convertir. Solo se reserva el rango para
    # que las listas desplegables lleguen hasta abajo.
    ultima = fila + filas_libres

    validacion(hoja, rango('paises'), 'D4:D%d' % ultima)
    validacion(hoja, rango('tipos'), 'E4:E%d' % ultima)
    validacion(hoja, rango('donde'), 'F4:F%d' % ultima)
    validacion(hoja, rango('severidad'), 'K4:K%d' % ultima)
    validacion(hoja, rango('si_no'), 'L4:L%d' % ultima)
    validacion(hoja, rango('estado_juan'), 'N4:N%d' % ultima)
    return hoja


# ------------------------------------------------------ datos legales país ----

CAMPOS_PAIS = [
    ('entidad_legal', 'Sociedad de Rappi', 'Es quien figura como la plataforma en el documento, con su nombre completo.', 'Alta'),
    ('id_fiscal', 'Identificación fiscal de Rappi', 'Va en la misma frase de la sociedad.', 'Alta'),
    ('domicilio_legal', 'Ciudad de domicilio', 'Sale junto a la sociedad y en la cláusula de ley aplicable.', 'Media'),
    ('ley_aplicable', 'Ley aplicable', 'Cierra el documento: bajo qué leyes se interpreta.', 'Alta'),
    ('jurisdiccion', 'Jurisdicción', 'Los jueces o tribunales que conocen de una controversia.', 'Alta'),
    ('entidad_vigilancia', 'Autoridad de consumo', 'Es la autoridad a la que el documento dice que el usuario puede acudir.', 'Alta'),
    ('autoridad_datos', 'Autoridad de datos personales', 'Aparece en la cláusula de tratamiento de datos.', 'Alta'),
    ('texto_derechos_datos', 'Frase de los derechos del titular', 'Es la lista de derechos sobre los datos personales, dentro de la cláusula de datos. Tiene que salir de la ley de tu país, no de una traducción de la frase colombiana.', 'Alta'),
    ('politica_privacidad', 'Nombre de la política de privacidad', 'El nombre exacto del documento publicado, como se cita en el texto.', 'Alta'),
    ('edad_minima', 'Edad mínima para participar', 'Va en la cláusula de participantes, en letras y en números.', 'Alta'),
    ('tipo_id_organizador', 'Cómo se identifica al Organizador', 'El tipo de identificación de la tienda o la marca que organiza.', 'Media'),
    ('texto_persona', 'Cómo se nombra a la persona', 'Persona natural, persona física, persona humana.', 'Media'),
    ('moneda_nombre', 'Cómo se escribe la moneda', 'Acompaña a todos los montos del documento.', 'Media'),
    ('moneda_simbolo', 'Símbolo de la moneda', 'El símbolo que va pegado a la cifra.', 'Media'),
    ('separador_miles', 'Separador de miles', 'Si mil se escribe 1.000 o 1,000.', 'Media'),
    ('termino_dias', 'Cómo se llaman los días', 'Días calendario, días corridos, días naturales, días hábiles.', 'Media'),
    ('texto_hora_local', 'Cómo se nombra la hora local', 'Cierra las fechas de inicio y de fin de la campaña.', 'Media'),
    ('url_tc_plataforma', 'Enlace a los T&C de la plataforma', 'El documento lo cita y el usuario hace clic. Confirma que abra y que sea el vigente.', 'Alta'),
    ('url_privacidad', 'Enlace a la política de privacidad', 'Mismo caso: confirma que abra y que sea el vigente.', 'Alta'),
    ('url_tc_creditos', 'Enlace a los T&C de Créditos', 'Lo citan todas las campañas que dan Créditos.', 'Alta'),
    ('url_acuerdo_datos', 'Acuerdo de transferencia de datos', 'Es el documento que firma la tienda o la marca cuando recibe datos de los ganadores. Si tu país no tiene uno, dilo.', 'Media'),
]

CAMPOS_SOLO_CO = [
    ('travel_entidad', 'Sociedad de Rappi Travel', 'Solo para las campañas de Rappi Travel.', 'Alta'),
    ('travel_id_fiscal', 'Identificación fiscal de Rappi Travel', 'Solo para las campañas de Rappi Travel.', 'Alta'),
    ('travel_rnt', 'Registro Nacional de Turismo', 'Va impreso en los documentos de Travel.', 'Alta'),
    ('turbo_entidad', 'Sociedad detrás del botón Turbo', 'Sale en la campaña de 15 minutos o Cashback.', 'Alta'),
    ('url_tc_turbo', 'Enlace a los T&C del botón Turbo', 'Lo cita la campaña de 15 minutos o Cashback.', 'Media'),
]

NOTAS_CAMPO = {
    ('BR', 'texto_derechos_datos'):
        'Va vacío a propósito: el modelo aprobado de Brasil no lleva cláusula de datos. Dinos si '
        'la LGPD obliga a incluirla.',
}


def pestana_datos_pais(libro, paises):
    cabeceras = [
        'País', 'Dato', 'Para qué sirve y dónde sale', 'Prioridad',
        'Lo que dice hoy la herramienta', '¿Está bien?',
        'El valor correcto, escrito exacto', 'Norma, enlace o fuente',
        'Quién lo confirma', '[Juan] Estado', '[Juan] Nota',
    ]
    anchos = [14, 32, 52, 12, 56, 20, 52, 36, 22, 14, 30]
    hoja = hoja_nueva(
        libro, '3 Datos de tu país',
        '3. Los datos legales de tu país',
        'Filtra por tu país en la columna A. Estos valores entran tal cual en cada documento que '
        'se genere allá, así que un error aquí sale en todos. Lo verde es lo que dice hoy la '
        'herramienta: no lo edites, corrige en la columna G.',
        cabeceras, anchos)

    fila = 4
    for pais in paises:
        campos = list(CAMPOS_PAIS)
        if pais['codigo'] == 'CO':
            campos += CAMPOS_SOLO_CO
        for clave, etiqueta, explicacion, prioridad in campos:
            valor = pais.get(clave, '')
            nota = NOTAS_CAMPO.get((pais['codigo'], clave))
            if nota:
                explicacion = explicacion + ' ' + nota
            if valor == '':
                valor = '(vacío)'
            escribir_fila(hoja, fila, [pais['nombre'], etiqueta, explicacion, prioridad, valor])
            hoja.cell(row=fila, column=5).fill = FONDO_PREFIJADO
            fila += 1

    validacion(hoja, rango('esta_bien'), 'F4:F%d' % (fila - 1))
    validacion(hoja, rango('estado_juan'), 'J4:J%d' % (fila - 1))
    return hoja


# ---------------------------------------------------- preguntas de fondo -----

PREGUNTAS = [
    ('¿En tu país hay que registrar, autorizar o notificar una promoción ante alguna autoridad antes de publicarla?',
     'Si tu país pide un trámite previo, la herramienta tiene que avisarlo en pantalla antes de que el comercial publique. Dinos cuál es la autoridad, con cuánta antelación y para qué tipos de campaña.',
     'Alta'),
    ('¿Tu ley exige una duración mínima, un aviso previo para modificar o terminar la campaña, o un tiempo mínimo entre publicar el T&C y arrancar?',
     'Hoy el documento dice que la campaña puede terminar cuando se agote el presupuesto. Si en tu país eso no se puede, o necesita un aviso previo, hay que cambiar la frase.',
     'Alta'),
    ('¿Los Créditos pueden tener fecha de vencimiento en tu país? ¿Hay una vigencia mínima obligatoria?',
     'La herramienta deja poner la vigencia en días o en una fecha. Si tu ley tiene un piso, lo ponemos como mínimo y no dejamos bajar de ahí.',
     'Alta'),
    ('Rappi nunca figura como Organizador: organiza la tienda o la marca que fondea, y Rappi es intermediario y responsable del tratamiento de datos. ¿Esa construcción se sostiene en tu país?',
     'Es la regla que sostiene todo el producto. Si en tu país la autoridad considera organizador a quien opera la mecánica, necesito saberlo ya y con la norma en la mano.',
     'Alta'),
    ('¿Se puede poner tope por pedido, presupuesto total que termina la campaña antes de la fecha y límite de pedidos por usuario, sin que eso sea publicidad engañosa o cláusula abusiva?',
     'Son las tres limitaciones que más usa el negocio. Si alguna no pasa en tu país, la quitamos de las preguntas para ese país.',
     'Alta'),
    ('¿Qué categorías de producto no pueden entrar en una campaña de cashback o de descuento en tu país?',
     'Alcohol, tabaco y nicotina, medicamentos, fórmulas infantiles, apuestas, armas, cualquier otra. En Colombia cerramos medicamentos, fitoterapéuticos y nicotina. Dinos cuáles cierra tu país para bloquearlas también.',
     'Alta'),
    ('Si un T&C generado sale con un error y ya está publicado, ¿cómo se arregla en tu país?',
     'Hay tres caminos posibles: corregir la misma página, publicar una modificación con aviso, o respetar lo publicado hasta que la campaña termine. La herramienta ya sabe generar una versión nueva con nota de modificación al final, pero necesito saber si eso basta.',
     'Alta'),
    ('¿Basta con publicar el T&C en la página de Squarespace y enlazarlo desde la pieza publicitaria?',
     'Si tu país exige otro canal, un formato específico o que el documento siga publicado un tiempo mínimo después de que termine la campaña, lo escribimos en los pasos de publicación.',
     'Media'),
    ('¿El documento tiene que estar en el idioma oficial de tu país y con la moneda local, sin excepción?',
     'Hoy Brasil genera en portugués y el resto en español. Si tu país acepta o exige algo distinto, lo ajustamos.',
     'Media'),
    ('¿Qué advertencia o leyenda obligatoria debe llevar la pieza publicitaria en tu país?',
     'La respuesta larga va en la pestaña 5, con el texto exacto. Aquí solo dinos si tu país tiene leyendas obligatorias y cuántas.',
     'Media'),
]


def pestana_preguntas(libro, paises):
    cabeceras = [
        'País', 'Número', 'Pregunta', 'Por qué te la hago', 'Prioridad',
        'Respuesta', 'Detalle o el texto exacto que debería ir',
        'Norma o fuente', 'Quién responde', '[Juan] Estado', '[Juan] Nota',
    ]
    anchos = [14, 9, 58, 52, 12, 16, 56, 32, 22, 14, 30]
    hoja = hoja_nueva(
        libro, '4 Preguntas de Legal',
        '4. Diez preguntas de fondo sobre tu país',
        'Filtra por tu país y responde las diez. Son las que no puedo contestar desde Colombia. '
        'Si una no aplica, escribe "no aplica" y sigue.',
        cabeceras, anchos)

    fila = 4
    for pais in paises:
        for numero, (pregunta, porque, prioridad) in enumerate(PREGUNTAS, start=1):
            fila = escribir_fila(
                hoja, fila,
                [pais['nombre'], numero, pregunta, porque, prioridad])

    validacion(hoja, rango('respuesta'), 'F4:F%d' % (fila - 1))
    validacion(hoja, rango('estado_juan'), 'J4:J%d' % (fila - 1))
    return hoja


# --------------------------------------------------------------- leyendas ----

def pestana_leyendas(libro, leyendas, filas_libres=40):
    cabeceras = [
        'País', 'Identificador', 'Categoría', 'En qué campañas',
        'Cómo se usa hoy en la herramienta',
        'El texto que se imprime hoy', 'Largo máximo',
        '¿Está bien?', 'El texto correcto, escrito exacto',
        'Norma que lo obliga', '¿Falta alguna leyenda obligatoria en tu país?',
        '[Juan] Estado', '[Juan] Nota',
    ]
    anchos = [14, 32, 26, 26, 34, 62, 12, 20, 56, 32, 42, 14, 28]
    hoja = hoja_nueva(
        libro, '5 Leyendas',
        '5. Las leyendas de las piezas publicitarias',
        'Son las advertencias que van impresas en la pieza, no en el T&C. Filtra por tu país. Si '
        'falta una que tu ley obliga, agrégala en las filas vacías del final con tu país en la '
        'columna A.',
        cabeceras, anchos)

    fila = 4
    for leyenda in leyendas:
        se_elige = DONDE_SE_ELIGE.get(
            leyenda.get('se_elige_en', ''), leyenda.get('se_elige_en', ''))
        tipos = leyenda.get('tipo_campana', '')
        tipos = 'Todas' if tipos == 'todas' else tipos
        fila = escribir_fila(hoja, fila, [
            paises_en_letra(leyenda.get('pais', '')),
            leyenda.get('id', leyenda['archivo']),
            leyenda.get('categoria', ''),
            tipos,
            se_elige,
            leyenda.get('texto', ''),
            leyenda.get('largo_max', ''),
        ])
        hoja.cell(row=fila - 1, column=6).fill = FONDO_PREFIJADO

    # Las filas vacias no se pintan una por una: el archivo se hincha y Google
    # les pone su propio formato al convertir. Solo se reserva el rango para
    # que las listas desplegables lleguen hasta abajo.
    ultima = fila + filas_libres

    validacion(hoja, rango('esta_bien'), 'H4:H%d' % ultima)
    validacion(hoja, rango('estado_juan'), 'L4:L%d' % ultima)
    return hoja


# ------------------------------------------------- accesos a Squarespace -----

ACCESOS_HOY = [
    ['Colombia', 'Camilo Patiño', 'ivan.patino@rappi.com', 'Legal Colombia', 'Legal',
     'Publicar T&C de cashback', 'Sí', '', 'Juan Camilo Gallego', 'Ya lo tenía'],
    ['Colombia', 'Valery Medina', 'ext-v.reina@rappi.com', 'Legal Colombia', 'Legal',
     'Publicar T&C de cashback', 'Sí', '', 'Juan Camilo Gallego', 'Ya lo tenía'],
]


def pestana_squarespace(libro, filas_libres=120):
    cabeceras = [
        'País', 'Nombre completo', 'Correo corporativo', 'Equipo o vertical',
        'Cargo', '¿Para qué necesita el acceso?',
        '¿Ya tiene usuario de Squarespace?',
        'Cuántos T&C publica al mes, más o menos',
        'Abogado del país que lo aprueba',
        '[Maca] Estado del acceso', '[Maca] Fecha en que quedó creado', '[Juan] Nota',
    ]
    anchos = [14, 26, 30, 24, 16, 28, 22, 24, 26, 20, 20, 30]
    hoja = hoja_nueva(
        libro, '6 Accesos Squarespace',
        '6. Quién necesita usuario de Squarespace',
        'Squarespace es donde se publica el T&C, así que dar acceso es dar la llave del sitio '
        'legal. Por eso cada nombre tiene que venir aprobado por el abogado del país, columna I. '
        'Una fila por persona. Maca crea los usuarios y marca las columnas grises.',
        cabeceras, anchos)

    fila = 4
    for acceso in ACCESOS_HOY:
        fila = escribir_fila(hoja, fila, acceso, fondo=FONDO_PREFIJADO)

    # Las filas vacias no se pintan una por una: el archivo se hincha y Google
    # les pone su propio formato al convertir. Solo se reserva el rango para
    # que las listas desplegables lleguen hasta abajo.
    ultima = fila + filas_libres

    validacion(hoja, rango('paises'), 'A4:A%d' % ultima)
    validacion(hoja, rango('para_que'), 'F4:F%d' % ultima)
    validacion(hoja, rango('si_no'), 'G4:G%d' % ultima)
    validacion(hoja, rango('estado_acceso'), 'J4:J%d' % ultima)
    return hoja


# -------------------------------------------------- accesos a RappiMind ------

def pestana_accesos_rappimind(libro, filas_libres=80):
    cabeceras = [
        'País', 'Nombre completo', 'Correo corporativo', 'Rol en la herramienta',
        '¿Va a probar esta semana?', 'Notas',
        '[Juan] Estado', '[Juan] Fecha',
    ]
    anchos = [14, 26, 30, 22, 22, 40, 16, 14]
    hoja = hoja_nueva(
        libro, '7 Accesos RappiMind',
        '7. Quién más de tu equipo quiere probar',
        'Escribe aquí a quien quieras sumar a las pruebas y yo lo doy de alta. El rol "legal" ve '
        'su país y puede generar y leer todo. El rol "comercial" es el que tendrá el KAM cuando '
        'abramos. Nadie ve la herramienta si no está en esta lista.',
        cabeceras, anchos)

    ultima = 4 + filas_libres

    validacion(hoja, rango('paises'), 'A4:A%d' % ultima)
    validacion(hoja, rango('rol_rappimind'), 'D4:D%d' % ultima)
    validacion(hoja, rango('si_no'), 'E4:E%d' % ultima)
    validacion(hoja, rango('estado_alta'), 'G4:G%d' % ultima)
    return hoja


# ------------------------------------------------------ cobertura futura -----

def pestana_cobertura(libro, paises):
    cabeceras = [
        'País', 'Lo que ya puede generar hoy',
        'La campaña que más te piden y la herramienta no hace',
        'La segunda que más te piden',
        '¿Tienen T&C generales publicados para descuentos o cupones?',
        'Enlaces de esos T&C generales',
        'Cuántas solicitudes de T&C recibes al mes, más o menos',
        'Quién te las pide más',
        'Lo que quieras contarme', '[Juan] Nota',
    ]
    anchos = [14, 46, 40, 40, 28, 46, 26, 24, 40, 28]
    hoja = hoja_nueva(
        libro, '8 Cobertura',
        '8. Qué falta por cubrir en tu país',
        'Esto define qué construimos después. Si la herramienta no te quita trabajo real, no '
        'sirve: dime qué es lo que de verdad te llega todas las semanas.',
        cabeceras, anchos)

    fila = 4
    for pais in paises:
        tipos = ', '.join(NOMBRES_TIPOS.get(m, m) for m in pais.get('modelos', []))
        generales = 'Sí, 14 publicados' if pais['codigo'] == 'CO' else ''
        fila = escribir_fila(hoja, fila, [pais['nombre'], tipos, '', '', generales])
        hoja.cell(row=fila - 1, column=2).fill = FONDO_PREFIJADO

    validacion(hoja, rango('si_no'), 'E4:E%d' % (fila - 1))
    return hoja


# ---------------------------------------------------------------- resumen ----

def pestana_resumen(libro, paises):
    cabeceras = [
        'País', 'Hallazgos', 'Bloqueantes', 'Serios', 'Menores', 'Sugerencias',
        'Datos por cambiar', 'Preguntas respondidas de 10',
        'Accesos pedidos a Squarespace', 'Cómo va el país',
    ]
    anchos = [16, 12, 14, 12, 12, 14, 18, 24, 26, 20]
    hoja = hoja_nueva(
        libro, '9 Resumen',
        '9. El resumen, que se llena solo',
        'No escribas aquí. Todo sale de las otras pestañas con fórmulas. Sirve para ver de un '
        'vistazo quién va cerrando y dónde está el riesgo.',
        cabeceras, anchos, corte='B4')

    hall = "'2 Hallazgos'"
    datos = "'3 Datos de tu país'"
    preg = "'4 Preguntas de Legal'"
    sqsp = "'6 Accesos Squarespace'"
    chk = "'1 Checklist'"

    fila = 4
    for indice, pais in enumerate(paises):
        nombre = pais['nombre']
        fila_checklist = 4 + indice
        valores = [
            nombre,
            '=COUNTIF(%s!$D$6:$D$500,$A%d)' % (hall, fila),
            '=COUNTIFS(%s!$D$6:$D$500,$A%d,%s!$K$6:$K$500,"Bloqueante")' % (hall, fila, hall),
            '=COUNTIFS(%s!$D$6:$D$500,$A%d,%s!$K$6:$K$500,"Serio")' % (hall, fila, hall),
            '=COUNTIFS(%s!$D$6:$D$500,$A%d,%s!$K$6:$K$500,"Menor")' % (hall, fila, hall),
            '=COUNTIFS(%s!$D$6:$D$500,$A%d,%s!$K$6:$K$500,"Sugerencia")' % (hall, fila, hall),
            '=COUNTIFS(%s!$A$4:$A$300,$A%d,%s!$F$4:$F$300,"Hay que cambiarlo")' % (datos, fila, datos),
            '=COUNTIFS(%s!$A$4:$A$120,$A%d,%s!$F$4:$F$120,"<>")' % (preg, fila, preg),
            '=COUNTIF(%s!$A$4:$A$200,$A%d)' % (sqsp, fila),
            '=%s!$I$%d' % (chk, fila_checklist),
        ]
        fila = escribir_fila(hoja, fila, valores)

    total = fila
    escribir_fila(hoja, total, [
        'Todos',
        '=SUM(B4:B%d)' % (total - 1),
        '=SUM(C4:C%d)' % (total - 1),
        '=SUM(D4:D%d)' % (total - 1),
        '=SUM(E4:E%d)' % (total - 1),
        '=SUM(F4:F%d)' % (total - 1),
        '=SUM(G4:G%d)' % (total - 1),
        '=SUM(H4:H%d)' % (total - 1),
        '=SUM(I4:I%d)' % (total - 1),
        '=COUNTIF(%s!$I$4:$I$%d,"Terminado")&" de 9 países cerrados"' % (chk, 3 + len(paises)),
    ], fuente=FUERTE)
    return hoja


# ------------------------------------------------------------------ main -----

def main():
    salida = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        RAIZ, 'RappiMind_revision_legal_global.xlsx')

    paises = leer_paises()
    leyendas = leer_leyendas()

    libro = Workbook()
    libro.remove(libro.active)

    pestana_leeme(libro)
    pestana_checklist(libro, paises)
    pestana_hallazgos(libro)
    pestana_datos_pais(libro, paises)
    pestana_preguntas(libro, paises)
    pestana_leyendas(libro, leyendas)
    pestana_squarespace(libro)
    pestana_accesos_rappimind(libro)
    pestana_cobertura(libro, paises)
    pestana_resumen(libro, paises)
    pestana_listas(libro)

    libro.save(salida)
    print('Hoja escrita en %s' % salida)
    print('%d países, %d leyendas, %d pestañas' % (
        len(paises), len(leyendas), len(libro.sheetnames)))


if __name__ == '__main__':
    main()
